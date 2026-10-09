"""Isolated auth/roster acceptance: mocked WeCom, real local transactions and signed JWTs."""
import asyncio
import hashlib
import io
import os
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import patch
from urllib.parse import parse_qs, urlparse

sandbox = tempfile.TemporaryDirectory(prefix="aftersales-auth-")
os.environ.update(DATABASE_URL_OVERRIDE="sqlite+aiosqlite:///" + sandbox.name + "/test.db",
                  BACKGROUND_JOBS="false", LOGIN_MODE="password", ENVIRONMENT="development")
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
from httpx import AsyncClient, ASGITransport
from sqlalchemy import select, update, delete
from app.bootstrap import initialize, seed
from app.main import app
from app.core.config import settings, Settings
from app.core.database import async_session, engine
from app.core.security import create_access_token, verify_password
from app.models.employee import Employee
from app.models.employee_role import EmployeeRole
from app.models.field_service import FieldLoginState, FieldWecomIdentity, FieldSchemaVersion
from app.services import wecom_auth

passed = []


async def main():
    await initialize()
    await seed(True)
    async with async_session() as db:
        db.add(Employee(name='旧状态人员', employee_no='LEGACY-STATUS', phone='13800000021', status='停用', is_active=False, hashed_password='unchanged'))
        await db.execute(delete(FieldSchemaVersion).where(FieldSchemaVersion.version == 6))
        await db.commit()
    await initialize()
    async with async_session() as db:
        legacy = await db.scalar(select(Employee).where(Employee.employee_no=='LEGACY-STATUS'))
        assert legacy.status=='在职' and not legacy.is_active and legacy.hashed_password=='unchanged'
    passed.append('legacy labels normalize without unlocking accounts or changing credentials')
    app.state.limiter.enabled = False
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        async def req(method, path, token=None, status=200, **kw):
            headers = {"Authorization": "Bearer " + token} if token else {}
            r = await client.request(method, "/api/v1" + path, headers=headers, **kw)
            assert r.status_code == status, (method, path, r.status_code, r.text[:600])
            return r.json() if "json" in r.headers.get("content-type", "") else r.content
        admin = (await req("POST", "/auth/management-login", json={"phone": "19900000001", "password": "Demo-AfterSales-2026!"}))["access_token"]
        staff = (await req("POST", "/auth/login", json={"phone": "19900000003", "password": "Demo-AfterSales-2026!"}))["access_token"]
        await req("GET", "/auth/management-session", admin)
        await req("GET", "/auth/management-session", status=401)
        for path in ["/auth/management-session", "/attendance/rules", "/field/dashboard", "/field/people"]:
            await req("GET", path, staff, status=403)
        await req("GET", "/auth/me", staff)
        await req("GET", "/attendance/my/runtime", staff)
        passed.append("mobile engineer sessions cannot access management session, rules, overview or roster")
        assert (await req("GET", "/auth/policy"))["password_enabled"]
        await req("POST", "/auth/wecom/start", status=503, json={"client":"web","environment":"web","challenge":"a"*64})
        passed.append("local password policy and unconfigured WeCom fail closed")
        files = {"file": ("roster.csv", "姓名,手机号,工号,岗位,在职状态,企微账号\n测试工程师,13800000001,,售后工程师,在职,\n重复,13800000001,,,,\n错误,not-phone,,,,\n离职人员,13800000002,,,离职,\n".encode())}
        await req("POST", "/field/roster/import", staff, status=403, files=files)
        await req("GET", "/field/roster/template", staff, status=403)
        preview = await req("POST", "/field/roster/import", admin, files=files)
        assert (preview["success"], preview["failed"]) == (2, 2)
        assert not any(p["phone"] == "13800000001" for p in await req("GET", "/field/people", admin))
        applied = await req("POST", "/field/roster/import?preview=false", admin, files=files)
        assert (applied["success"], applied["failed"]) == (2, 2)
        new = (await req("POST", "/auth/login", json={"phone":"13800000001","password":"585858"}))["access_token"]
        await req("POST", "/auth/management-login", status=403, json={"phone":"13800000001","password":"585858"})
        await req("POST", "/auth/login", status=401, json={"phone":"13800000002","password":"585858"})
        async with async_session() as db:
            employee = await db.scalar(select(Employee).where(Employee.phone == "13800000001"))
            employee_id = employee.id
            assert verify_password("585858", employee.hashed_password)
            old_hash = employee.hashed_password
            db.add(EmployeeRole(employee_id=employee_id, role_name="hr"))
            await db.commit()
        repeated = await req("POST", "/field/roster/import?preview=false", admin, files=files)
        assert repeated["success"] == 2
        async with async_session() as db:
            assert (await db.get(Employee, employee_id)).hashed_password == old_hash
            assert await db.scalar(select(EmployeeRole.id).where(EmployeeRole.employee_id == employee_id, EmployeeRole.role_name == "hr"))
        passed.append("preview makes no writes; phone accounts, initial password, duplicates, disabled admission and RBAC")
        conflict = await req("POST", "/field/roster/import", admin, files={"file":("roster.csv", "姓名,手机号\n冒名,13800000001\n".encode())})
        assert conflict["failed"] == 1
        from openpyxl import Workbook
        wb = Workbook(); ws=wb.active; ws.append(["姓名","手机号"]); ws.append(["表格员工",13800000005]); buf=io.BytesIO(); wb.save(buf)
        assert (await req("POST", "/field/roster/import", admin, files={"file":("test.xlsx",buf.getvalue())}))["success"] == 1
        passed.append("existing credentials/roles survive repeat imports; mismatched names reject; xlsx numeric phone parses")
        # The download and UI share the exact reference headers; never ship real people.
        from openpyxl import load_workbook
        expected_headers = ['姓名','职位','所属部门','所属公司','当前所在项目','入职日期','离职日期','售后组长','联系电话','衣服尺码','身份证号','年龄','是否进入企微','入职天数','公司电脑','夏季工服','反光马甲','电脑建议','公司在职状态']
        template = await req('GET', '/field/roster/template', admin)
        workbook = load_workbook(io.BytesIO(template))
        sheet = workbook.active
        assert sheet.max_row == 2 and sheet.max_column == 19
        assert [c.value for c in sheet[1]] == expected_headers
        assert sheet.freeze_panes == 'A2'
        columns = await req('GET', '/field/roster/columns', admin)
        assert [c['label'] for c in columns] == expected_headers
        assert (await req('POST', '/field/roster/import', admin, files={'file': ('template.xlsx', template)}))['failed'] == 1
        values = ['名册测试人员','售后工程师','测试部门','测试公司','登记项目','2026/09/01',None,'测试组长','13800000009','XL','110101199601010010',30,'是',29,'已配发','2件','1件','需要更换','在职']
        for cell_, value in zip(sheet[2], values):
            cell_.value = value
        buf = io.BytesIO(); workbook.save(buf)
        roster_file = {'file': ('roster.xlsx', buf.getvalue())}
        assert (await req('POST', '/field/roster/import', admin, files=roster_file))['success'] == 1
        assert not any(p['phone']=='13800000009' for p in await req('GET', '/field/people', admin))
        assert (await req('POST', '/field/roster/import?preview=false', admin, files=roster_file))['success'] == 1
        roster_person = next(p for p in await req('GET', '/field/people', admin) if p['phone']=='13800000009')
        values[5] = '2026-09-01'
        for column, value in zip(columns, values):
            assert roster_person[column['key']] == value, (column, roster_person[column['key']], value)
        assert not roster_person['wecom_userid']  # Registration is not an identity binding.
        await req('PUT', f"/field/people/{roster_person['id']}", admin, json={**roster_person, 'clothing_size':'L', 'is_active':False})
        changed = next(p for p in await req('GET', '/field/people', admin) if p['id']==roster_person['id'])
        assert changed['clothing_size']=='L' and changed['status']=='在职' and not changed['is_active']
        await req('PUT', f"/field/people/{changed['id']}", admin, status=422, json={**changed, 'age':-1})
        await req('PUT', f"/field/people/{changed['id']}", admin, status=422, json={**changed, 'leave_date':'2026-08-01'})
        sheet['K2'] = '=1+1'
        buf = io.BytesIO(); workbook.save(buf)
        invalid = await req('POST', '/field/roster/import?preview=false', admin, files={'file':('invalid.xlsx',buf.getvalue())})
        assert invalid['failed']==1 and '身份证号' in invalid['rows'][0]['error']
        assert next(p for p in await req('GET','/field/people',admin) if p['id']==changed['id'])['clothing_size']=='L'
        added = await req('POST', '/field/people', admin, json={'name':'手工名册人员','phone':'13800000008','company_computer':'未配发'})
        assert next(p for p in await req('GET','/field/people',admin) if p['id']==added['id'])['employee_no']=='AF-13800000008'
        manual_token = (await req('POST','/auth/login',json={'phone':'13800000008','password':'585858'}))['access_token']
        await req('POST','/auth/change-password',manual_token,status=400,json={'old_password':'wrong','new_password':'Employee-New-2026!'})
        await req('POST','/auth/change-password',manual_token,json={'old_password':'585858','new_password':'Employee-New-2026!'})
        await req('POST','/auth/login',status=401,json={'phone':'13800000008','password':'585858'})
        await req('POST','/auth/login',json={'phone':'13800000008','password':'Employee-New-2026!'})
        await req('POST','/auth/management-login',status=403,json={'phone':'13800000008','password':'Employee-New-2026!'})
        manual = next(p for p in await req('GET','/field/people',admin) if p['id']==added['id'])
        for obsolete in ['待入职','停用']:
            await req('PUT',f"/field/people/{manual['id']}",admin,status=422,json={**manual,'status':obsolete})
            invalid_status = await req('POST','/field/roster/import',admin,files={'file':('status.csv',f'姓名,联系电话,公司在职状态\n手工名册人员,13800000008,{obsolete}\n'.encode())})
            assert invalid_status['failed']==1
        await req('PUT',f"/field/people/{manual['id']}",admin,json={**manual,'status':'试用期'})
        assert next(p for p in await req('GET','/field/people',admin) if p['id']==manual['id'])['status']=='试用'
        await req('GET','/auth/me',manual_token)
        await req('PUT',f"/field/people/{manual['id']}",admin,json={**manual,'status':'离职','is_active':True})
        await req('GET','/auth/me',manual_token,status=403)
        await req('POST','/auth/login',status=401,json={'phone':'13800000008','password':'585858'})
        retained = next(p for p in await req('GET','/field/people',admin) if p['id']==manual['id'])
        assert retained['status']=='离职' and not retained['is_active'] and retained['company_computer']=='未配发'
        assert manual['role']=='employee'
        lead = await req('POST','/field/people',admin,json={'name':'管理权限验收','phone':'13800000029','role':'hr'})
        lead_token = (await req('POST','/auth/management-login',json={'phone':'13800000029','password':'585858'}))['access_token']
        for path in ['/auth/management-session','/field/people','/field/dashboard','/attendance/rules']:
            await req('GET',path,lead_token)
        await req('POST','/field/people',lead_token,json={'name':'普通员工默认权限验收','phone':'13800000028'})
        await req('POST','/field/people',lead_token,status=403,json={'name':'禁止授权管理员','phone':'13800000027','role':'admin'})
        lead_row = next(p for p in await req('GET','/field/people',admin) if p['id']==lead['id'])
        await req('PUT',f"/field/people/{lead['id']}",admin,json={**lead_row,'role':'employee'})
        await req('GET','/auth/management-session',lead_token,status=403)
        passed.append('manual entry opens phone account; exactly three employment states; probation works; departure revokes login and preserves profile')
        passed.append('exact 19-column template; one fictional sample; all fields round-trip; preview isolation; company status independent of account switch; validation')
        settings.WECHAT_CORP_ID="ww-test-corp"
        settings.WECHAT_AGENT_ID="1000001"
        settings.WECHAT_SECRET="test-only"
        settings.WECHAT_WEB_REDIRECT_URI="https://service.example.test/login"
        settings.WECHAT_MOBILE_REDIRECT_URI="https://service.example.test/mobile/login"
        settings.LOGIN_MODE="hybrid"
        assert settings.wecom_ready
        people=await req("GET", "/field/people",admin)
        person=next(p for p in people if p["id"]==employee_id)
        await req("PUT", f"/field/people/{employee_id}",admin,json={**person,"wecom_userid":"engineer-1"})
        other=next(p for p in people if p["phone"]=="19900000003")
        await req("PUT",f"/field/people/{other['id']}",admin,status=409,json={**other,"wecom_userid":"engineer-1"})
        passed.append("explicit CorpID/userid bindings cannot be stolen")
        async def start(kind="mobile",env="in_app"):
            verifier="b"*64
            data=await req("POST","/auth/wecom/start",json={"client":kind,"environment":env,"challenge":hashlib.sha256(verifier.encode()).hexdigest()})
            url=urlparse(data["authorize_url"])
            args=parse_qs(url.query)
            assert args["appid"]==[settings.WECHAT_CORP_ID]
            assert args["redirect_uri"]==[settings.WECHAT_MOBILE_REDIRECT_URI if kind=="mobile" else settings.WECHAT_WEB_REDIRECT_URI]
            return {"client":kind,"verifier":verifier,"state":data["state"],"code":"one-time-code"}
        user_identity = {"userid":"engineer-1"}
        member_state = {"userid":"engineer-1","status":1}
        application = {"agentid":1000001,"close":0,"allow_userinfos":{"user":[{"userid":"engineer-1"},{"userid":"unknown"}]}}
        async def provider(path, params):
            if path=="gettoken": return {"access_token":"mock-provider-token","expires_in":7200}
            if path=="auth/getuserinfo": return user_identity
            if path=="user/get": return member_state
            if path=="agent/get": return application
            if path=="department/list": return {"department":[{"id":2},{"id":3}]}
            if path=="tag/get": return {"userlist":[{"userid":"engineer-1"}],"partylist":[]}
            raise AssertionError(path)
        with patch.object(wecom_auth,"provider_get",provider):
            attempt=await start()
            await req("POST","/auth/wecom/complete",status=400,json={**attempt,"verifier":"c"*64})
            await req("POST","/auth/wecom/complete",status=400,json={**attempt,"client":"web"})
            result=await req("POST","/auth/wecom/complete",json=attempt)
            wecom_token=result["access_token"]
            await req("POST","/auth/wecom/complete",status=400,json=attempt)
            await req("GET","/auth/me",wecom_token)
            passed.append("mobile OAuth callback verifier/client isolation, successful identity exchange and replay rejection")
            original_application = application
            application = {"agentid":1000001,"close":1}
            await req("POST","/auth/wecom/complete",status=403,json=await start())
            application["close"]=0
            await req("POST","/auth/wecom/complete",status=403,json=await start())
            application["allow_partys"]={"partyid":[2]}
            member_state["department"]=[3]
            await req("POST","/auth/wecom/complete",json=await start())
            application={"agentid":1000001,"close":0,"allow_tags":{"tagid":[9]}}
            await req("POST","/auth/wecom/complete",json=await start())
            application = original_application
            passed.append("application disabled/out-of-scope denial; visible members, departments, child departments and tags")
            expired=await start()
            async with async_session() as db:
                await db.execute(update(FieldLoginState).where(FieldLoginState.digest==wecom_auth.digest(expired['state'])).values(expires_at=datetime.utcnow()-timedelta(seconds=1)))
                await db.commit()
            await req("POST","/auth/wecom/complete",status=400,json=expired)
            user_identity={"openid":"external"}
            nonmember=await start()
            await req("POST","/auth/wecom/complete",status=403,json=nonmember)
            await req("POST","/auth/wecom/complete",status=400,json=nonmember)
            user_identity={"userid":"anotherCorp/worker"}
            await req("POST","/auth/wecom/complete",status=403,json=await start())
            user_identity={"userid":"unknown"};member_state={"userid":"unknown","status":1}
            await req("POST","/auth/wecom/complete",status=403,json=await start())
            user_identity={"userid":"engineer-1"};member_state={"userid":"engineer-1","status":2}
            await req("POST","/auth/wecom/complete",status=403,json=await start())
            member_state["status"]=1
            passed.append("expired states, non-enterprise members, unrostered and disabled WeCom members denied")
            web_attempt=await start('web','web')
            await req("POST","/auth/wecom/complete",json=web_attempt)
            await req("PUT",f"/field/people/{employee_id}",admin,json={**person,"role":"employee","wecom_userid":"engineer-1"})
            await req("POST","/auth/wecom/complete",status=403,json=await start('web','web'))
            await req("GET", "/auth/management-session", new, status=403)
            await req("GET", "/auth/me", new)
            passed.append("desktop WeCom login and existing sessions recheck management-role boundary")
            settings.LOGIN_MODE="wecom_only"
            assert not (await req("GET","/auth/policy"))["password_enabled"]
            for endpoint in ["login","management-login"]:
                await req("POST","/auth/"+endpoint,status=403,json={"phone":"19900000001","password":"Demo-AfterSales-2026!"})
            await req("GET","/auth/me",admin,status=401)
            await req("GET","/field/people",admin,status=401)
            await req("GET","/auth/me",wecom_token)
            passed.append("WeCom-only mode closes both password endpoints and rejects previously issued password tokens")
            async with async_session() as db:
                binding=await db.scalar(select(FieldWecomIdentity).where(FieldWecomIdentity.employee_id==employee_id))
                binding.userid="rebound"
                await db.commit()
            await req("GET","/auth/me",wecom_token,status=401)
            passed.append("rebinding revokes WeCom sessions immediately")
        configured = Settings(_env_file=None,ENVIRONMENT="production",SECRET_KEY="x"*40,LOGIN_MODE="password")
        assert configured.LOGIN_MODE == "password"
        passed.append("production supports the roster phone/password mode with an explicit signing secret")
    await engine.dispose()
    print(f"PASS {len(passed)} auth/roster acceptance groups (WeCom provider mocked)")
    for item in passed: print(" -",item)

if __name__=="__main__": asyncio.run(main())
