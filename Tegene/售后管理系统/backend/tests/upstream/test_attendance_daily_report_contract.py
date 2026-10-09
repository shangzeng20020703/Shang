"""
日报汇总与月报联动契约测试
覆盖:
1) 日报查询字段与导出字段来自同一套列定义
2) 列选择解析逻辑可按 key 子集导出
3) 打卡后月报增量刷新逻辑（锁定月份跳过）
"""

from __future__ import annotations

import asyncio
from io import BytesIO
from datetime import date, datetime, time, timezone
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.schemas.attendance import MonthSummaryGenerateRequest
from app.services.attendance import AttendancePunchTimeRecordService, AttendanceRecordService, AttendanceReportService, MonthSummaryService
from openpyxl import load_workbook


class TestDailyReportColumnContract:
    def test_daily_overview_column_specs_are_unique_and_export_consistent(self):
        specs = AttendanceReportService.DAILY_OVERVIEW_COLUMN_SPECS
        keys = [item["key"] for item in specs]
        titles = [item["title"] for item in specs]
        assert len(keys) == len(set(keys))
        assert AttendanceReportService.DAILY_OVERVIEW_COLUMNS == titles

    def test_daily_detail_column_specs_are_unique_and_export_consistent(self):
        specs = AttendanceReportService.DAILY_DETAIL_COLUMN_SPECS
        keys = [item["key"] for item in specs]
        titles = [item["title"] for item in specs]
        assert len(keys) == len(set(keys))
        assert AttendanceReportService.DAILY_DETAIL_COLUMNS == titles

    def test_resolve_column_specs_uses_requested_subset_and_ignores_unknown_keys(self):
        specs = AttendanceReportService.DAILY_DETAIL_COLUMN_SPECS
        resolved = AttendanceReportService._resolve_column_specs(
            specs,
            ["date_label", "employee_name", "not_exists"],
        )
        assert [item["key"] for item in resolved] == ["date_label", "employee_name"]

    def test_clock_count_is_limited_to_clock_in_and_clock_out_slots(self):
        clock_in_only = SimpleNamespace(
            clock_in_time=datetime(2026, 4, 24, 9, 0, tzinfo=timezone.utc),
            clock_out_time=None,
        )
        full_day = SimpleNamespace(
            clock_in_time=datetime(2026, 4, 24, 9, 0, tzinfo=timezone.utc),
            clock_out_time=datetime(2026, 4, 24, 18, 0, tzinfo=timezone.utc),
        )

        assert AttendanceReportService._attendance_clock_count(None) == 0
        assert AttendanceReportService._attendance_clock_count(clock_in_only) == 1
        assert AttendanceReportService._attendance_clock_count(full_day) == 2

    def test_standard_work_hours_uses_schedule_times_and_rest_periods(self):
        rule = SimpleNamespace(
            clock_in_time=time(9, 0),
            clock_out_time=time(18, 0),
            work_hours_per_day=Decimal("7"),
        )
        runtime = {
            "expected_in": time(9, 0),
            "expected_out": time(18, 0),
            "extra": {"lunch_as_working": False},
            "rest_periods": [(time(12, 0), time(13, 0))],
        }

        assert AttendanceReportService._standard_work_hours_from_runtime(rule, runtime, date(2026, 4, 24)) == Decimal("8")

    def test_free_rule_uses_hours_limit_instead_of_whole_day_punch_window(self):
        rule = SimpleNamespace(
            work_hour_type="flexible", clock_in_time=None, clock_out_time=None,
            work_hours_per_day=Decimal("12"),
            extra_config={"rule_type": "free", "free_work_hours_mode": "limit:12"},
        )
        runtime = {"expected_in": time(0, 0), "expected_out": time(23, 59), "extra": rule.extra_config}
        assert AttendanceReportService._standard_work_hours_from_runtime(rule, runtime, date(2026, 10, 9)) == Decimal("12")
        rule.extra_config["free_work_hours_mode"] = "unlimited"
        assert AttendanceReportService._standard_work_hours_from_runtime(rule, runtime, date(2026, 10, 9)) == Decimal("0")

    def test_standard_work_hours_rounds_final_result_to_integer(self):
        rule = SimpleNamespace(
            clock_in_time=time(10, 26),
            clock_out_time=time(10, 30),
            work_hours_per_day=Decimal("8"),
        )
        runtime = {
            "expected_in": time(10, 26),
            "expected_out": time(10, 30),
            "extra": {"lunch_as_working": False},
            "rest_periods": [],
        }

        assert AttendanceReportService._standard_work_hours_from_runtime(rule, runtime, date(2026, 4, 24)) == Decimal("0")

    def test_standard_work_hours_does_not_use_unset_segment_rest_from_legacy_lunch_fields(self):
        rule = SimpleNamespace(
            clock_in_time=time(10, 0),
            clock_out_time=time(10, 30),
            work_hours_per_day=Decimal("8"),
        )
        runtime = {
            "expected_in": time(10, 0),
            "expected_out": time(10, 30),
            "extra": {"lunch_as_working": False, "lunch_start": "12:00", "lunch_end": "13:00"},
            "rest_periods": [],
        }

        assert AttendanceReportService._standard_work_hours_from_runtime(rule, runtime, date(2026, 4, 29)) == Decimal("1")


class TestMonthlyReportColumnContract:
    def test_monthly_overview_column_specs_are_unique_and_export_consistent(self):
        specs = AttendanceReportService.MONTHLY_OVERVIEW_COLUMN_SPECS
        keys = [item["key"] for item in specs]
        titles = [item["title"] for item in specs]
        assert len(keys) == len(set(keys))
        assert AttendanceReportService.MONTHLY_OVERVIEW_COLUMNS == titles
        assert "punch_correction_count" in keys

    def test_monthly_overview_columns_match_template_order(self):
        assert AttendanceReportService.MONTHLY_OVERVIEW_COLUMNS == [
            "姓名",
            "账号",
            "所属规则",
            "部门",
            "职务",
            "应出勤天数(天)",
            "实际出勤天数(天)",
            "休息天数(天)",
            "正常天数(天)",
            "异常天数(天)",
            "标准工作时长(小时)",
            "实际工作时长(小时)",
            "异常合计(次)",
            "迟到次数(次)",
            "迟到时长(分钟)",
            "早退次数(次)",
            "早退时长(分钟)",
            "旷工次数(次)",
            "旷工时长(分钟)",
            "缺卡次数(次)",
            "地点异常(次)",
            "设备异常(次)",
            "补卡次数(次)",
            "审批打卡次数(次)",
            "外勤次数(次)",
            "外出(小时)",
            "出差(天)",
            "年假(天)",
            "婚假(天)",
            "产休假(小时)",
            "陪产假(天)",
            "病假(天)",
            "调休(小时)",
            "事假(天)",
            "病假（长期）(天)",
            "产检假(天)",
            "丧假(天)",
            "病假（通用）(天)",
            "加班时长(小时)",
            "工作日加班时长(小时)",
            "工作日加班计为调休(小时)",
            "工作日加班计为加班费(小时)",
            "休息日加班时长(小时)",
            "休息日加班计为调休(小时)",
            "休息日加班计为加班费(小时)",
            "节假日加班时长(小时)",
            "节假日加班计为调休(小时)",
            "节假日加班计为加班费(小时)",
        ]

    def test_monthly_detail_base_column_specs_are_unique_and_export_consistent(self):
        specs = AttendanceReportService.MONTHLY_DETAIL_BASE_COLUMN_SPECS
        keys = [item["key"] for item in specs]
        titles = [item["title"] for item in specs]
        assert len(keys) == len(set(keys))
        assert AttendanceReportService.MONTHLY_DETAIL_BASE_COLUMNS == titles

    def test_monthly_day_specs_follow_selected_date_range(self):
        specs = AttendanceReportService._month_day_column_specs(
            date(2026, 4, 20),
            date(2026, 4, 22),
        )
        assert [item["key"] for item in specs] == ["day_20260420", "day_20260421", "day_20260422"]
        assert specs[0]["title"] == "20\n星期一"
        assert specs[-1]["date"] == "2026-04-22"

    def test_monthly_status_filter_ignores_invalid_placeholder_values(self):
        row = {"_normal_days": 0, "_abnormal_days": 1, "_status_tokens": {"迟到"}}
        assert AttendanceReportService._monthly_status_matches(row, None) is True
        assert AttendanceReportService._monthly_status_matches(row, "") is True
        assert AttendanceReportService._monthly_status_matches(row, "打卡状态") is True

    def test_monthly_day_detail_appends_effect_annotations(self):
        record = SimpleNamespace(
            clock_in_time=datetime(2026, 4, 10, 1, 0, tzinfo=timezone.utc),
            clock_out_time=datetime(2026, 4, 10, 10, 30, tzinfo=timezone.utc),
        )
        outside_effect = SimpleNamespace(
            effect_type="outside",
            minutes=90,
            start_at=datetime(2026, 4, 10, 9, 0, tzinfo=timezone.utc),
            end_at=datetime(2026, 4, 10, 10, 30, tzinfo=timezone.utc),
            work_date=date(2026, 4, 10),
            payload_json={},
        )

        detail = AttendanceReportService._monthly_day_detail(
            record,
            expected_work_day=True,
            status_text="正常",
            day_effects=[outside_effect],
            leave_items=[],
            standard_minutes=480,
        )

        assert "正常（外出）" in detail
        assert "外出1.5小时" in detail

    def test_export_wecom_monthly_report_uses_single_template_sheet(self):
        overview_rows = [
            {
                "employee_name": "AAA",
                "account": "AAA",
                "rule_name": "上下班打卡",
                "department_name": "天下先智创机器人/人事部",
                "position": "--",
                "punch_correction_count": "1次",
            }
        ]
        detail_rows = [
            {
                "employee_name": "AAA",
                "account": "AAA",
                "rule_name": "上下班打卡",
                "department_name": "天下先智创机器人/人事部",
                "position": "--",
                "day_20260401": "正常 ;",
            }
        ]
        day_columns = [{"key": "day_20260401", "title": "1\n星期三", "date": "2026-04-01"}]

        with patch.object(
            AttendanceReportService,
            "_load_monthly_report_rows",
            new=AsyncMock(return_value=(overview_rows, detail_rows, 1, day_columns)),
        ):
            data = asyncio.run(
                AttendanceReportService.export_wecom_monthly_report(
                    None,
                    start_date=date(2026, 4, 1),
                    end_date=date(2026, 4, 1),
                )
            )

        workbook = load_workbook(BytesIO(data))
        assert workbook.sheetnames == ["上下班打卡_月报"]
        sheet = workbook["上下班打卡_月报"]
        assert sheet.freeze_panes == "B5"
        assert sheet["A1"].value == "上下班打卡_月报"
        assert "统计时间:04-01 ～ 04-01" in sheet["A2"].value
        header_values = [cell.value for cell in sheet[4]]
        assert "补卡次数\n(次)" in header_values
        assert "1\n星期三" in header_values


class TestPunchTimeRecordContract:
    def test_punch_time_date_columns_follow_wecom_calendar_header(self):
        columns = AttendancePunchTimeRecordService.date_columns(
            date(2026, 4, 1),
            date(2026, 4, 3),
        )
        assert [item["title"] for item in columns] == ["1\n星期三", "2\n星期四", "3\n星期五"]
        assert columns[0]["date"] == "2026-04-01"

    def test_punch_time_record_range_is_limited_to_month_window(self):
        try:
            AttendancePunchTimeRecordService.validate_month_window(date(2026, 4, 1), date(2026, 5, 2))
        except ValueError as exc:
            assert "最大仅支持按月查询" in str(exc)
        else:
            raise AssertionError("expected monthly window validation to fail")


class _FakeDB:
    def __init__(self, locked_count: int):
        self.locked_count = locked_count

    async def scalar(self, *_args, **_kwargs):
        return self.locked_count


class TestMonthSummaryRefreshBridge:
    def test_refresh_month_summary_generates_when_month_not_locked(self):
        fake_db = _FakeDB(locked_count=0)
        with patch.object(MonthSummaryService, "generate", new=AsyncMock()) as mocked_generate:
            asyncio.run(
                AttendanceRecordService._refresh_month_summary_for_date(
                    fake_db,
                    employee_id=23,
                    target_date=date(2026, 4, 22),
                )
            )
            assert mocked_generate.await_count == 1
            args = mocked_generate.await_args.args
            assert args[0] is fake_db
            payload = args[1]
            assert isinstance(payload, MonthSummaryGenerateRequest)
            assert payload.year == 2026
            assert payload.month == 4
            assert payload.employee_ids == [23]

    def test_refresh_month_summary_skips_when_month_locked(self):
        fake_db = _FakeDB(locked_count=1)
        with patch.object(MonthSummaryService, "generate", new=AsyncMock()) as mocked_generate:
            asyncio.run(
                AttendanceRecordService._refresh_month_summary_for_date(
                    fake_db,
                    employee_id=23,
                    target_date=date(2026, 4, 22),
                )
            )
            assert mocked_generate.await_count == 0
