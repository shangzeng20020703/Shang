"""Preview-first roster import. Existing roles, passwords and disabled access survive imports."""
import csv
import io
import json
import re
import zipfile

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from pydantic import ValidationError

from app.core.config import settings
from app.core.security import DEFAULT_INITIAL_PASSWORD, get_password_hash
from app.models.employee import Employee
from app.models.employee_role import EmployeeRole
from app.models.field_service import FieldWecomIdentity, FieldRosterProfile
from app.schemas.field_roster import ROSTER_COLUMNS, RosterFields, RosterDetails
from app.models.audit_log import AuditLog
from app.services.field_service import require_admin, roles
from app.services.wecom_auth import ACTIVE_STATUSES, bind_member

HEADERS = {
    "姓名": "name", "员工姓名": "name", "员工名称": "name",
    "手机号": "phone", "手机号码": "phone", "联系电话": "phone", "手机": "phone",
    "工号": "employee_no", "员工编号": "employee_no", "员工工号": "employee_no",
    "岗位": "position", "职位": "position", "岗位名称": "position",
    "状态": "status", "员工状态": "status", "在职状态": "status",
    "企微账号": "wecom_userid", "企业微信账号": "wecom_userid", "userid": "wecom_userid",
    "企微userid": "wecom_userid", "企业微信userid": "wecom_userid",
}
HEADERS.update({label: key for key, label, _ in ROSTER_COLUMNS})
LIMIT_ROWS = 1000


async def save_roster_fields(db, employee, values):
    details = {key: value for key, value in values.items() if key in RosterDetails.model_fields}
    for key, value in values.items():
        if key in RosterFields.model_fields and key not in RosterDetails.model_fields:
            setattr(employee, key, value)
    if details:
        profile = await db.get(FieldRosterProfile, employee.id)
        if profile is None:
            profile = FieldRosterProfile(employee_id=employee.id, details={})
            db.add(profile)
        profile.details = {**profile.details, **details}


def cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_file(content, filename):
    if not content or len(content) > 5 * 1024 * 1024:
        raise HTTPException(400, "文件不能为空，且最多 5 MB")
    name = (filename or "").lower()
    try:
        if name.endswith(".xlsx"):
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError("解压后的文件过大")
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                ws = wb.worksheets[0]
                if ws.max_column is None or ws.max_row is None:
                    ws.calculate_dimension(force=True)
                if ws.max_column > 100 or ws.max_row > LIMIT_ROWS + 20:
                    raise ValueError("最多 1000 条数据、100 列")
                rows = list(ws.iter_rows(values_only=True))
            finally:
                wb.close()
        elif name.endswith(".xls"):
            import xlrd
            book = xlrd.open_workbook(file_contents=content)
            sheet = book.sheet_by_index(0)
            if sheet.nrows > LIMIT_ROWS + 20 or sheet.ncols > 100:
                raise ValueError("最多 1000 条数据、100 列")
            rows = [sheet.row_values(i) for i in range(sheet.nrows)]
        elif name.endswith(".csv"):
            try:
                decoded = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                decoded = content.decode("gb18030")
            rows = list(csv.reader(io.StringIO(decoded)))
            if len(rows) > LIMIT_ROWS + 20 or any(len(r) > 100 for r in rows):
                raise ValueError("最多 1000 条数据、100 列")
        else:
            raise ValueError("请选择 xlsx、xls 或 csv 花名册")
        for index, row in enumerate(rows[:10]):
            headers = [HEADERS.get(re.sub(r"[\s*＊]", "", cell(v)).lower()) for v in row]
            if "name" in headers and "phone" in headers:
                known = [h for h in headers if h]
                if len(set(known)) != len(known):
                    raise ValueError("存在重复的姓名、手机号或其他已识别列")
                parsed = [(i + index + 2, {h: cell(v) for h, v in zip(headers, r) if h})
                          for i, r in enumerate(rows[index + 1:]) if any(cell(v) for v in r)]
                if len(parsed) > LIMIT_ROWS:
                    raise ValueError("每次最多导入 1000 名员工")
                return parsed
        raise ValueError("前十行内未找到‘姓名’和‘手机号’表头")
    except (ValueError, KeyError, OSError, zipfile.BadZipFile) as exc:
        raise HTTPException(400, str(exc)) from None
    except Exception:
        raise HTTPException(400, "无法读取文件，请重新导出标准 Excel 或 CSV") from None


async def import_roster(db, operator, content, filename, dry_run=True):
    await require_admin(db, operator)
    parsed = parse_file(content, filename)
    report, seen_phones, seen_numbers, seen_members = [], set(), set(), set()
    for line, source in parsed:
        phone = re.sub(r"[\s-]", "", source.get("phone", ""))
        if phone.startswith("+86"):
            phone = phone[3:]
        item = {"row": line, "name": source.get("name", ""), "phone": phone,
                "employee_no": source.get("employee_no", ""), "action": "", "error": ""}
        try:
            if not item["name"] or len(item["name"]) > 50 or item["name"].startswith("="):
                raise ValueError("姓名必填，最多 50 字，不接受公式")
            if not re.fullmatch(r"1[3-9]\d{9}", phone):
                raise ValueError("手机号须为 11 位有效号码；不接受公式")
            if phone in seen_phones:
                raise ValueError("文件内手机号重复")
            seen_phones.add(phone)
            employee = await db.scalar(select(Employee).where(Employee.phone == phone))
            if employee and (employee.is_superuser or "admin" in await roles(db, employee)):
                raise ValueError("系统管理员账号请在人员编辑中维护")
            placeholder = employee and employee.name in {"售后负责人1", "售后负责人2"}
            if employee and employee.name != item["name"] and not placeholder:
                raise ValueError("手机号已存在但姓名不同，请先人工核对")
            number = item["employee_no"] or (employee.employee_no if employee else "AF-" + phone)
            item["employee_no"] = number
            if len(number) > 30 or number.startswith("=") or number in seen_numbers:
                raise ValueError("工号过长、含公式或文件内重复")
            seen_numbers.add(number)
            number_owner = await db.scalar(select(Employee.id).where(Employee.employee_no == number))
            if number_owner and (not employee or number_owner != employee.id):
                raise ValueError("工号已被其他员工使用")
            if len(source.get("position", "")) > 100:
                raise ValueError("岗位最多 100 字")
            fields = {key: value for key, value in source.items() if key in RosterFields.model_fields}
            # Validate the combined dates even when an older file omits one of them.
            dates = {key: getattr(employee, key) for key in ("hire_date", "leave_date")} if employee else {}
            validated = RosterFields(**{**dates, **fields})
            fields = validated.model_dump(include=set(fields))
            status = source.get("status") or (employee.status if employee else "在职")
            if status == "试用期":
                status = "试用"
            if status not in ACTIVE_STATUSES | {"离职"}:
                raise ValueError("公司在职状态只能为在职、试用期或离职")
            if employee and (employee.id == operator.id) and status not in ACTIVE_STATUSES:
                raise ValueError("不能通过花名册停用自己的账号")
            member = source.get("wecom_userid", "").lower()
            if member:
                if not settings.WECHAT_CORP_ID:
                    raise ValueError("请先配置企业 CorpID，或暂时留空企微账号列")
                if len(member.encode("utf-8")) > 64 or "/" in member or member.startswith("=") or member in seen_members:
                    raise ValueError("企微账号过长、含公式或文件内重复")
                seen_members.add(member)
                bound = await db.scalar(select(FieldWecomIdentity).where(
                    FieldWecomIdentity.corp_id == settings.WECHAT_CORP_ID,
                    FieldWecomIdentity.userid == member))
                if bound and (not employee or bound.employee_id != employee.id):
                    raise ValueError("企微账号已绑定其他员工")
                existing_binding = await db.scalar(select(FieldWecomIdentity).where(
                    FieldWecomIdentity.employee_id == employee.id)) if employee else None
                if existing_binding and (existing_binding.userid != member or
                                         existing_binding.corp_id != settings.WECHAT_CORP_ID):
                    raise ValueError("已有企微绑定不同，请在人员编辑中核对后修改")
            item.update(action="更新资料" if employee else "新建账号", status=status,
                        wecom_userid=member, position=source.get("position", ""),
                        note="保留原密码、角色和停用状态" if employee else "初始角色：售后人员")
            item.update(fields)
            if not dry_run:
                async with db.begin_nested():
                    if not employee:
                        employee = Employee(phone=phone, is_field_staff=True,
                                            hashed_password=get_password_hash(DEFAULT_INITIAL_PASSWORD),
                                            is_active=status in ACTIVE_STATUSES)
                        db.add(employee)
                    employee.name, employee.employee_no, employee.status = item["name"], number, status
                    if source.get("position"):
                        employee.position = source["position"]
                    await db.flush()
                    await save_roster_fields(db, employee, fields)
                    if status not in ACTIVE_STATUSES:
                        employee.is_active = False
                    await db.flush()
                    if not employee.is_active:
                        from app.services.field_membership import stop_person
                        await stop_person(db, employee.id)
                    if item["action"] == "新建账号":
                        db.add(EmployeeRole(employee_id=employee.id, role_name="employee", granted_by=operator.id))
                    if member:
                        await bind_member(db, employee.id, member)
                    await db.flush()
        except ValidationError as exc:
            labels = {key: label for key, label, _ in ROSTER_COLUMNS}
            item["error"] = "；".join(f"{labels.get(e['loc'][0], e['loc'][0]) if e['loc'] else '日期'}：{e['msg']}" for e in exc.errors())
        except (ValueError, HTTPException) as exc:
            item["error"] = exc.detail if isinstance(exc, HTTPException) else str(exc)
        except IntegrityError:
            item["error"] = "工号、手机号或企微账号发生并发冲突，请刷新后重试"
        report.append(item)
    result = {"preview": dry_run, "total": len(report),
              "success": sum(not r["error"] for r in report),
              "failed": sum(bool(r["error"]) for r in report), "rows": report}
    if not dry_run:
        db.add(AuditLog(operator_id=operator.id, operator_name=operator.name, action="import",
                       module="field", resource_type="roster",
                       after_data=json.dumps({k: result[k] for k in ("total", "success", "failed")})))
    return result
